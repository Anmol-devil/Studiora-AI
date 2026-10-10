
import fitz
import docx

from pathlib import Path
from langchain_text_splitters import RecursiveCharacterTextSplitter


def extract_text(file_path):
    pages = []
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        with fitz.open(file_path) as pdf:
            for page_num, page in enumerate(pdf):
                pages.append({
                    "page": page_num + 1,
                    "text": page.get_text("text")
                })

    elif extension == ".docx":
        document = docx.Document(file_path)
        content = []

        # Extract paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                content.append(paragraph.text)

        # Extract tables
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                content.append(" | ".join(cells))

        pages.append({
            "page": 1,
            "text": "\n".join(content)
        })

    elif extension == ".txt":
        with open(file_path, "r", encoding="utf-8-sig") as file:
            pages.append({
                "page": 1,
                "text": file.read()
            })

    elif extension == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. "
            "Save the file as .docx and upload it again."
        )

    else:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )

    # Validate extracted text
    total_chars = sum(
        len(page["text"].strip()) for page in pages
    )

    print(f"File: {file_path}")
    print(f"Extracted pages: {len(pages)}")
    print(f"Extracted characters: {total_chars}")

    if total_chars == 0:
        raise ValueError(
            "No text was extracted from this file."
        )

    return pages


def create_chunks(text):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=100
    )
    return splitter.split_text(text)

